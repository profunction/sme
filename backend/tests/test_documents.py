import io
import time
import zipfile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader

from backend import documents, store, tasks
from backend.tests.test_app import AUTH

PAYLOAD = {
    "company": "NightBrew",
    "customer": "Customer <safe>",
    "reference": "Q-100",
    "tax_percent": "6",
    "lines": [{"description": "Tea", "quantity": "3", "unit_price": "10.05"}],
}


def wait(client, identifier):
    for _ in range(1000):
        row = next(
            t
            for t in client.get("/tasks", headers=AUTH).json()
            if t["id"] == identifier
        )
        if row["status"] not in ("queued", "running", "cancel_requested"):
            return row
        time.sleep(0.01)
    raise AssertionError("Task did not finish")


def test_decimal_totals_and_generated_pdf():
    model = documents.BusinessDocument(**PAYLOAD)
    data, summary = documents.generate(model)
    assert summary == {
        "subtotal": "30.15",
        "tax": "1.81",
        "total": "31.96",
        "currency": "MYR",
    }
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    assert "31.96" in text and "Customer <safe>" in text
    assert (
        documents.totals(
            documents.BusinessDocument(
                **{
                    **PAYLOAD,
                    "lines": [
                        {
                            "description": "round",
                            "quantity": "1.005",
                            "unit_price": "1.00",
                        }
                    ],
                }
            )
        )[1]["total"]
        == "1.07"
    )


def test_generate_download_extract(client):
    request = client.post(
        "/tasks",
        headers=AUTH,
        json={"tool": "document.generate", "payload": {**PAYLOAD, "format": "docx"}},
    )
    assert request.status_code == 201
    task = wait(client, request.json()["id"])
    assert task["status"] == "completed"
    identifier = task["result"]["file"]["id"]
    downloaded = client.get("/files/" + identifier, headers=AUTH)
    assert "NightBrew" in "\n".join(
        p.text for p in Document(io.BytesIO(downloaded.content)).paragraphs
    )
    task_id = client.post(
        "/tasks",
        headers=AUTH,
        json={"tool": "document.extract", "payload": {"file_id": identifier}},
    ).json()["id"]
    extracted = wait(client, task_id)
    assert extracted["result"]["untrusted_content"] is True
    assert "NightBrew" in extracted["result"]["text"]


def test_excel_no_formula_execution(client):
    book = Workbook()
    sheet = book.active
    sheet.append(["Item", "Quantity"])
    sheet.append(["Tea", 3])
    sheet.append(["Coffee", 5])
    sheet.append(["Formula", "=SUM(B2:B3)"])
    output = io.BytesIO()
    book.save(output)
    uploaded = client.post(
        "/files?filename=stock.xlsx", headers=AUTH, content=output.getvalue()
    )
    assert uploaded.status_code == 201
    identifier = client.post(
        "/tasks",
        headers=AUTH,
        json={"tool": "excel.analyze", "payload": {"file_id": uploaded.json()["id"]}},
    ).json()["id"]
    result = wait(client, identifier)
    assert result["status"] == "completed"
    assert result["result"]["sheets"][0]["numeric_columns"]["2"] == {
        "count": 2,
        "sum": "8",
    }


def test_delete_requires_single_use_approval(client):
    file = client.post(
        "/files?filename=data.csv", headers=AUTH, content=b"item,quantity\ntea,3"
    ).json()
    identifier = client.post(
        "/tasks",
        headers=AUTH,
        json={"tool": "file.delete", "payload": {"file_id": file["id"]}},
    ).json()["id"]
    assert client.get("/files/" + file["id"], headers=AUTH).status_code == 200
    assert (
        client.post("/tasks/" + identifier + "/approve", headers=AUTH).status_code
        == 200
    )
    assert wait(client, identifier)["status"] == "completed"
    assert client.get("/files/" + file["id"], headers=AUTH).status_code == 404
    assert (
        client.post("/tasks/" + identifier + "/approve", headers=AUTH).status_code
        == 409
    )


def test_cancel_pending_and_emergency_stop(client):
    file = client.post("/files?filename=data.csv", headers=AUTH, content=b"data").json()
    identifier = client.post(
        "/tasks",
        headers=AUTH,
        json={"tool": "file.delete", "payload": {"file_id": file["id"]}},
    ).json()["id"]
    assert (
        client.post("/tasks/" + identifier + "/cancel", headers=AUTH).status_code == 200
    )
    assert (
        client.post("/tasks/" + identifier + "/approve", headers=AUTH).status_code
        == 409
    )
    assert client.get("/files/" + file["id"], headers=AUTH).status_code == 200
    client.post("/control/stop", headers=AUTH)
    assert (
        client.post(
            "/tasks",
            headers=AUTH,
            json={"tool": "document.generate", "payload": PAYLOAD},
        ).status_code
        == 422
    )
    assert client.post("/chat", headers=AUTH, json={"message": "hi"}).status_code == 409
    assert client.post("/control/resume", headers=AUTH).status_code == 200


def test_rejects_paths_and_unsafe_archives(client):
    assert client.get("/files/not-a-file-id", headers=AUTH).status_code == 404
    assert (
        client.post(
            "/files?filename=program.exe", headers=AUTH, content=b"MZ"
        ).status_code
        == 422
    )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "x")
        z.writestr("bomb.xml", "x" * 1000000)
    assert (
        client.post(
            "/files?filename=data.xlsx", headers=AUTH, content=output.getvalue()
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/tasks",
            headers=AUTH,
            json={"tool": "shell.execute", "payload": {"command": "anything"}},
        ).status_code
        == 422
    )


def test_recovery_never_replays_external_actions(tmp_path):
    connect = lambda: store.open_db(tmp_path / "jobs.db")
    identifier, _ = tasks.create(connect, "document.generate", PAYLOAD)
    tasks.recover(connect)
    with connect() as db:
        assert (
            db.execute("SELECT status FROM tasks WHERE id=?", (identifier,)).fetchone()[
                0
            ]
            == "interrupted"
        )


def test_deadline_and_running_cancellation(tmp_path, monkeypatch):
    connect = lambda: store.open_db(tmp_path / "jobs.db")
    identifier, _ = tasks.create(connect, "document.generate", PAYLOAD)

    def slow(*args):
        tasks.cancel(connect, identifier)
        args[-1]()

    monkeypatch.setattr(documents, "run", slow)
    tasks.execute(connect, tmp_path / "jobs.db", identifier)
    with connect() as db:
        assert db.execute("SELECT status FROM tasks").fetchone()[0] == "cancelled"


def test_task_deadline_fails_without_creating_output(tmp_path, monkeypatch):
    connect = lambda: store.open_db(tmp_path / "jobs.db")
    identifier, _ = tasks.create(connect, "document.generate", PAYLOAD)
    monkeypatch.setattr(tasks, "TIMEOUT_SECONDS", -1)
    tasks.execute(connect, tmp_path / "jobs.db", identifier)
    with connect() as db:
        row = db.execute("SELECT status,error FROM tasks").fetchone()
        assert row["status"] == "failed"
        assert row["error"] == "Task exceeded deadline"
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 0
