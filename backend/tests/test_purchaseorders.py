"""Real local task/file boundaries, synthetic purchase orders, no vendor/network calls."""
import csv
import hashlib
import io

import pytest
from openpyxl import Workbook, load_workbook

from backend import app as service, files, purchaseorders, store, tasks, workforce
from backend.tests.test_app import AUTH
from backend.tests.test_documents import wait

CSV = b'sku,description,quantity,unit_price,line_total\nTEA,Tea,3,10.05,30.15\n'


def create(client, tool, payload, code=201):
    response = client.post('/tasks', headers=AUTH, json={'tool': tool, 'payload': payload})
    assert response.status_code == code, response.text
    return response.json()


def imported(client, content=CSV, name='purchase-order.csv'):
    response = client.post('/files?filename=' + name, headers=AUTH, content=content)
    assert response.status_code == 201
    result = wait(client, create(client, 'purchase_order.import', {'file_id': response.json()['id']})['id'])
    assert result['status'] == 'completed', result
    return result['result']['purchase_order']


def reviewed_payload(source, **overrides):
    return {**source, 'reference': 'PO-100', 'customer': 'Customer', 'currency': 'MYR',
            'tax_percent': '6', 'expected_total': '31.96', 'format': 'csv',
            'catalog': [{'sku': 'TEA', 'unit_price': '10.05'}], **overrides}


def test_complete_review_requires_approval_and_exports_exact_csv(client):
    source = imported(client)
    assert source['source_sha256'] == hashlib.sha256(CSV).hexdigest()
    payload = reviewed_payload(source, customer='  =HYPERLINK("bad")')
    review = wait(client, create(client, 'purchase_order.review', payload)['id'])['result']['po_review']
    assert review['valid'] and review['issues'] == []
    assert review['totals'] == {'subtotal': '30.15', 'tax': '1.81', 'total': '31.96', 'currency': 'MYR'}
    export = create(client, 'purchase_order.export', {**payload, 'review_hash': review['review_hash']})
    assert export['status'] == 'pending_approval'
    assert len(client.get('/files', headers=AUTH).json()) == 1
    tasks.execute(service.connect, service.DB, export['id'])
    assert len(client.get('/files', headers=AUTH).json()) == 1
    assert client.post('/tasks/' + export['id'] + '/approve', headers=AUTH).status_code == 200
    result = wait(client, export['id'])
    assert result['status'] == 'completed', result
    file = result['result']['file']
    data = client.get('/files/' + file['id'], headers=AUTH).content
    row = list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))[0]
    assert row['line_total'] == '30.15' and row['unit_price'] == '10.05'
    assert row['customer'].startswith("'") and '=HYPERLINK' in row['customer']
    assert result['result']['output_sha256'] == hashlib.sha256(data).hexdigest()
    assert client.post('/tasks/' + export['id'] + '/approve', headers=AUTH).status_code == 409
    assert any(event['event'] == 'purchase_order_exported' and 'review=' in event['detail']
               for event in client.get('/audit', headers=AUTH).json())


@pytest.mark.parametrize('change', ['rows', 'format', 'expected_total', 'reference'])
def test_changed_review_cannot_reuse_old_hash(client, change):
    payload = reviewed_payload(imported(client))
    report = purchaseorders.validate(purchaseorders.Review(**payload))
    altered = {**payload, 'review_hash': report['review_hash']}
    if change == 'rows':
        altered['rows'] = [{**payload['rows'][0], 'description': 'Different goods'}]
    else:
        altered[change] = {'format': 'xlsx', 'expected_total': '35.00', 'reference': 'PO-OTHER'}[change]
    create(client, 'purchase_order.export', altered, 422)


def test_source_changed_after_review_fails_without_export(client):
    payload = reviewed_payload(imported(client))
    report = purchaseorders.validate(purchaseorders.Review(**payload))
    export = create(client, 'purchase_order.export', {**payload, 'review_hash': report['review_hash']})
    path, _ = files.locate(service.connect, service.DB, payload['file_id'])
    path.write_bytes(CSV.replace(b'Tea', b'Other goods'))
    client.post('/tasks/' + export['id'] + '/approve', headers=AUTH)
    assert wait(client, export['id'])['status'] == 'failed'
    assert len(client.get('/files', headers=AUTH).json()) == 1


@pytest.mark.parametrize('patch,field', [
    ({'quantity': '-3'}, 'quantity/unit_price/line_total'),
    ({'quantity': '0'}, 'quantity/unit_price/line_total'),
    ({'unit_price': 'NaN'}, 'quantity/unit_price/line_total'),
    ({'unit_price': '10.06', 'line_total': ''}, 'unit_price'),
    ({'sku': ''}, 'sku'),
    ({'description': ''}, 'description'),
    ({'line_total': '999.00'}, 'line_total'),
    ({'quantity': '=SUM(A1:A3)'}, 'quantity/unit_price/line_total'),
])
def test_discrepancies_block_export(client, patch, field):
    payload = reviewed_payload(imported(client))
    payload['rows'] = [{**payload['rows'][0], **patch}]
    result = wait(client, create(client, 'purchase_order.review', payload)['id'])['result']['po_review']
    assert not result['valid']
    assert any(issue['field'] == field for issue in result['issues'])
    create(client, 'purchase_order.export', {**payload, 'review_hash': result['review_hash']}, 422)


def test_xlsx_formula_is_not_evaluated_and_safe_export_has_no_formulas(client):
    book = Workbook()
    sheet = book.active
    sheet.append(['sku', 'description', 'quantity', 'unit_price', 'line_total'])
    sheet.append(['TEA', '=DANGEROUS()', '=1+2', '10.05', '30.15'])
    output = io.BytesIO()
    book.save(output)
    source = imported(client, output.getvalue(), 'purchase-order.xlsx')
    assert source['rows'][0]['quantity'] == '=1+2'
    payload = reviewed_payload(source, format='xlsx')
    assert not purchaseorders.validate(purchaseorders.Review(**payload))['valid']
    payload['rows'][0]['quantity'] = '3'
    report = purchaseorders.validate(purchaseorders.Review(**payload))
    assert report['valid']
    task = create(client, 'purchase_order.export', {**payload, 'review_hash': report['review_hash']})
    client.post('/tasks/' + task['id'] + '/approve', headers=AUTH)
    result = wait(client, task['id'])
    assert result['status'] == 'completed'
    data = client.get('/files/' + result['result']['file']['id'], headers=AUTH).content
    exported = load_workbook(io.BytesIO(data), data_only=False)
    assert all(cell.data_type != 'f' for row in exported.active for cell in row)
    assert exported.active['E2'].value.startswith("'")


def test_template_limits_and_duplicate_catalog_are_rejected(client):
    file = client.post('/files?filename=wrong.csv', headers=AUTH, content=b'SKU,price\nTEA,5\n').json()
    task = create(client, 'purchase_order.import', {'file_id': file['id']})
    assert wait(client, task['id'])['status'] == 'failed'
    payload = reviewed_payload(imported(client))
    payload['catalog'].append(payload['catalog'][0])
    report = purchaseorders.validate(purchaseorders.Review(**payload))
    assert not report['valid'] and any(issue['field'] == 'catalog' for issue in report['issues'])


def test_employee_cannot_disable_export_approval(client):
    payload = reviewed_payload(imported(client))
    report = purchaseorders.validate(purchaseorders.Review(**payload))
    employee = workforce.save(service.connect, workforce.Employee(name='PO worker', role='reviewer',
                               tools=['purchase_order.export'], require_approval=False))
    identifier, status = tasks.create(service.connect, 'purchase_order.export',
                        {**payload, 'review_hash': report['review_hash']}, employee['id'])
    assert status == 'pending_approval'
