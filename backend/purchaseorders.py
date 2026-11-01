"""Narrow, local PO review/export. No OCR, invented values or vendor compatibility.

Structured CSV/XLSX imports use exact column names. Catalog prices and expected
PO total must be supplied by the reviewing user. Export is an approved task.
"""
import csv
import hashlib
import io
import json
import re
from decimal import Decimal, ROUND_HALF_UP
from typing import Literal

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field

from backend import files, store


class Import(BaseModel):
    model_config = ConfigDict(extra="forbid")
    file_id: str = Field(pattern=r"^[a-f0-9]{32}$")


class Row(BaseModel):
    model_config = ConfigDict(extra="forbid")
    sku: str = Field(max_length=100)
    description: str = Field(max_length=200)
    quantity: str = Field(max_length=30)
    unit_price: str = Field(max_length=30)
    line_total: str = Field(default="", max_length=30)


class CatalogItem(BaseModel):
    model_config = ConfigDict(extra="forbid")
    sku: str = Field(min_length=1, max_length=100)
    unit_price: str = Field(min_length=1, max_length=30)


class Review(Import):
    source_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    reference: str = Field(min_length=1, max_length=80)
    customer: str = Field(min_length=1, max_length=150)
    currency: str = Field(default="MYR", pattern=r"^[A-Z]{3}$")
    tax_percent: str = Field(default="0", max_length=10)
    expected_total: str = Field(default="", max_length=30)
    format: Literal["csv", "xlsx"] = "csv"
    rows: list[Row] = Field(min_length=1, max_length=100)
    catalog: list[CatalogItem] = Field(min_length=1, max_length=100)


class Export(Review):
    review_hash: str = Field(pattern=r"^[a-f0-9]{64}$")


def money(value, places=2, maximum=Decimal("1000000000")):
    value = value.strip()
    if not re.fullmatch(r"\d+(?:\.\d{1," + str(places) + r"})?", value):
        raise ValueError("Use a nonnegative decimal without separators or formulas")
    number = Decimal(value)
    if number > maximum:
        raise ValueError("Value exceeds supported limit")
    return number


def fingerprint(document):
    data = document.model_dump(exclude={"review_hash"}, mode="json")
    return hashlib.sha256(json.dumps(data, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def validate(document):
    issues, catalog, seen = [], {}, set()

    def issue(row, field, message):
        issues.append({"row": row, "field": field, "message": message})

    for item in document.catalog:
        sku = item.sku.strip()
        try:
            if not sku or sku in catalog:
                raise ValueError("Catalog SKU is empty or duplicated")
            catalog[sku] = money(item.unit_price)
        except ValueError as exc:
            issue(None, "catalog", str(exc))
    subtotal, valid_numbers = Decimal("0"), True
    for index, row in enumerate(document.rows, 1):
        sku = row.sku.strip()
        if not sku or sku not in catalog:
            issue(index, "sku", "SKU must exist in the supplied reviewed catalog")
        if sku in seen:
            issue(index, "sku", "Duplicate SKU requires review; combine or correct the lines")
        seen.add(sku)
        if not row.description.strip():
            issue(index, "description", "Description requires confirmation")
        try:
            quantity = money(row.quantity, 3, Decimal("1000000"))
            if quantity <= 0:
                raise ValueError("Quantity must be greater than zero")
            price = money(row.unit_price)
            amount = (quantity * price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            subtotal += amount
            if sku in catalog and price != catalog[sku]:
                issue(index, "unit_price", "Price differs from the supplied catalog")
            if row.line_total.strip() and money(row.line_total, maximum=Decimal("1000000000000000")) != amount:
                issue(index, "line_total", "Reported line total differs from the calculated amount")
        except ValueError as exc:
            valid_numbers = False
            issue(index, "quantity/unit_price/line_total", str(exc))
    totals = None
    try:
        tax_percent = money(document.tax_percent, maximum=Decimal("100"))
        if valid_numbers:
            tax = (subtotal * tax_percent / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            total = subtotal + tax
            totals = {"subtotal": str(subtotal), "tax": str(tax), "total": str(total), "currency": document.currency}
            if money(document.expected_total, maximum=Decimal("200000000000000000")) != total:
                issue(None, "expected_total", "Reported PO total differs from the calculated total")
    except ValueError as exc:
        issue(None, "tax_percent/expected_total", str(exc))
    if not document.reference.strip() or not document.customer.strip():
        issue(None, "reference/customer", "Reference and customer require confirmation")
    return {"valid": not issues, "issues": issues, "totals": totals,
            "review_hash": fingerprint(document), "review": document.model_dump(exclude={"review_hash"}, mode="json")}


def prepare(tool, payload):
    if tool == "purchase_order.import":
        return Import.model_validate(payload).model_dump()
    model = Export.model_validate(payload) if tool == "purchase_order.export" else Review.model_validate(payload)
    if tool == "purchase_order.export":
        report = validate(model)
        if not report["valid"] or model.review_hash != report["review_hash"]:
            raise ValueError("Correct discrepancies and create a new reviewed export")
    return model.model_dump(mode="json")


def import_rows(path, checkpoint):
    required = ["sku", "description", "quantity", "unit_price"]
    book = None
    try:
        if path.suffix == ".csv":
            source = io.StringIO(path.read_text(encoding="utf-8-sig"), newline="")
            reader = csv.reader(source, strict=True)
        elif path.suffix == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            if len(book.worksheets) != 1:
                raise ValueError("PO import requires exactly one worksheet")
            reader = book.worksheets[0].iter_rows(values_only=True)
        else:
            raise ValueError("Structured PO import supports CSV/XLSX; PDF requires manual extraction/review")
        headings = [str(value).strip() if value is not None else "" for value in next(iter(reader), [])]
        if set(headings) not in (set(required), set(required + ["line_total"])) or len(set(headings)) != len(headings):
            raise ValueError("Use exact headers sku,description,quantity,unit_price and optional line_total")
        rows = []
        for values in reader:
            checkpoint()
            if all(value is None or value == "" for value in values):
                continue
            if len(rows) >= 100 or len(values) > len(headings):
                raise ValueError("Maximum 100 PO lines, with no extra columns")
            record = dict(zip(headings, ["" if value is None else str(value) for value in values]))
            rows.append(Row(**{name: record.get(name, "") for name in required + ["line_total"]}).model_dump())
        if not rows:
            raise ValueError("No PO lines found")
        return rows
    finally:
        if book:
            book.close()


def spreadsheet_text(value):
    # Prefix dangerous text for spreadsheet apps, including leading whitespace.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")) else value


def export_bytes(document):
    headings = ["reference", "customer", "currency", "sku", "description", "quantity", "unit_price", "line_total"]
    rows = []
    for row in document.rows:
        quantity, price = money(row.quantity, 3), money(row.unit_price)
        amount = (quantity * price).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
        rows.append([spreadsheet_text(value) for value in
                     [document.reference, document.customer, document.currency, row.sku, row.description]] +
                    [str(quantity), format(price, ".2f"), str(amount)])
    output = io.BytesIO()
    if document.format == "csv":
        text = io.StringIO(newline="")
        writer = csv.writer(text)
        writer.writerow(headings)
        writer.writerows(rows)
        return text.getvalue().encode("utf-8-sig")
    book = Workbook()
    sheet = book.active
    sheet.title = "Reviewed PO"
    sheet.append(headings)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # preserve exact decimals and never create formulas
    book.save(output)
    return output.getvalue()


def run(tool, payload, connect, db_path, checkpoint):
    checkpoint()
    path, _ = files.locate(connect, db_path, payload["file_id"])
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if tool == "purchase_order.import":
        return {"purchase_order": {"file_id": payload["file_id"], "source_sha256": source_hash,
                                   "rows": import_rows(path, checkpoint)},
                "note": "Structured template only. Review every value; no AI, OCR or ERP submission occurred."}
    document = Export.model_validate(payload) if tool == "purchase_order.export" else Review.model_validate(payload)
    if source_hash != document.source_sha256:
        raise ValueError("Source file changed; import and review it again")
    report = validate(document)
    if tool == "purchase_order.review":
        return {"po_review": report, "totals": report["totals"]}
    if not report["valid"] or document.review_hash != report["review_hash"]:
        raise ValueError("Reviewed values changed or contain discrepancies")
    data = export_bytes(document)
    checkpoint()
    file = files.save(connect, db_path, "reviewed-purchase-order." + document.format, data, trusted=True)
    with connect() as db:
        store.audit(db, "purchase_order_exported", file["id"], "source=" + source_hash + ";review=" + report["review_hash"])
    return {"file": file, "totals": report["totals"], "source_sha256": source_hash,
            "review_hash": report["review_hash"], "output_sha256": hashlib.sha256(data).hexdigest(),
            "note": "Neutral local template; vendor-specific mapping/connector acceptance remains required."}
